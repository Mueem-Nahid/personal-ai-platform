from __future__ import annotations

import subprocess
import sys
import types
from types import SimpleNamespace

import pytest

from renderers.base import RenderContext, RenderError, RendererUnavailableError
from renderers.docx_renderer import DocxRenderer
from renderers.html_renderer import HtmlRenderer
from renderers.latex_renderer import LatexRenderer
from renderers.pdf_import import import_pdf
from renderers.registry import capabilities, get_renderer, resolve_engine
from renderers.typst_renderer import TypstRenderer
from schemas.resume import ResumeContent, ResumeSection

CONTENT = ResumeContent(
    summary="Backend engineer with 6 years of experience & strong systems skills.",
    sections=[
        ResumeSection(
            name="Experience",
            items=[
                "Senior Engineer, Acme (2021-present) — Led #platform team; cut latency 45%",
                "Engineer, Globex (2018-2021) — Shipped payments service",
            ],
        ),
        ResumeSection(name="Skills", items=["Python, Go, PostgreSQL"]),
    ],
)

PROFILE = {
    "full_name": "Jane Doe",
    "title": "Senior Software Engineer",
    "email": "jane.doe@example.com",
    "phone": "+1 555 0100",
    "location": "Berlin, DE",
}

HTML_TEMPLATE = (
    "<html><head><title>x</title></head><body>"
    "<h1>{{ profile.full_name }}</h1>"
    "{% if content.summary %}<p>{{ content.summary }}</p>{% endif %}"
    "{% for s in content.sections %}<h2>{{ s.name }}</h2>"
    "{% for item in s.items %}<p>{{ item }}</p>{% endfor %}{% endfor %}"
    "</body></html>"
)

TYPST_TEMPLATE = (
    "= {{ profile.full_name|typst }}\n"
    "\n"
    "{{ profile.title|typst }}\n"
    "\n"
    "{% if content.summary %}{{ content.summary|typst }}{% endif %}\n"
    "\n"
    "{% for s in content.sections %}\n"
    "\n"
    "= {{ s.name|typst }}\n"
    "\n"
    "{% for item in s.items %}\n"
    "- {{ item|typst }}\n"
    "{% endfor %}"
    "{% endfor %}\n"
)

FIXTURE_TYP = r"""#set page(paper: "a4", margin: 1.5cm)
#text(size: 22pt)[Jane Doe]
#v(0.3em)
#text(size: 12pt)[Senior Software Engineer]
#v(0.3em)
#text(size: 9pt)[jane.doe\@example.com | +1 555 0100]
#v(0.8em)
#text(size: 14pt)[SUMMARY]
#v(0.2em)
#text(size: 10pt)[Backend engineer with many years of experience building systems.]
#v(0.6em)
#text(size: 14pt)[EXPERIENCE]
#v(0.2em)
#text(size: 10pt)[- Led platform team of six engineers]
#text(size: 10pt)[- Cut API latency by almost half]
#v(0.6em)
#text(size: 14pt)[SKILLS]
#v(0.2em)
#text(size: 10pt)[- Python, Go, PostgreSQL]
"""


def _ctx(**kwargs) -> RenderContext:
    defaults = dict(
        content=CONTENT,
        profile=dict(PROFILE),
        job={"title": "Staff Engineer", "company": "TechCorp"},
        template_source=TYPST_TEMPLATE,
    )
    defaults.update(kwargs)
    return RenderContext(**defaults)


@pytest.fixture
def weasyprint_stub(monkeypatch: pytest.MonkeyPatch):
    calls: list[str] = []

    class FakeHTML:
        def __init__(self, string: str, **kwargs: object) -> None:
            calls.append(string)

        def write_pdf(self) -> bytes:
            return b"%PDF-1.4 fake-weasyprint"

    module = types.ModuleType("weasyprint")
    module.HTML = FakeHTML  # type: ignore[attr-defined]
    monkeypatch.setitem(sys.modules, "weasyprint", module)
    return calls


class TestHtmlRenderer:
    async def test_renders_pdf_with_stub(self, weasyprint_stub: list[str]) -> None:
        renderer = HtmlRenderer()
        data = renderer.render(_ctx(template_source=HTML_TEMPLATE), "pdf")
        assert data.startswith(b"%PDF")
        html = weasyprint_stub[0]
        assert "Jane Doe" in html
        assert "Experience" in html
        assert "&amp;" in html  # autoescape on

    async def test_style_injection(self, weasyprint_stub: list[str]) -> None:
        renderer = HtmlRenderer()
        ctx = _ctx(template_source=HTML_TEMPLATE, styles_text="h1 { color: red }")
        renderer.render(ctx, "pdf")
        assert "<style>" in weasyprint_stub[0]
        assert "color: red" in weasyprint_stub[0]

    async def test_rejects_docx_output(self) -> None:
        with pytest.raises(RenderError, match="cannot produce 'docx'"):
            HtmlRenderer().render(_ctx(template_source=HTML_TEMPLATE), "docx")


class TestTypstRenderer:
    async def test_renders_real_pdf(self) -> None:
        data = TypstRenderer().render(_ctx(), "pdf")
        assert data.startswith(b"%PDF")
        assert len(data) > 1000

    async def test_rejects_docx_output(self) -> None:
        with pytest.raises(RenderError, match="cannot produce 'docx'"):
            TypstRenderer().render(_ctx(), "docx")


class TestLatexRenderer:
    async def test_disabled_by_default(self) -> None:
        with pytest.raises(RendererUnavailableError, match="disabled"):
            LatexRenderer().render(_ctx(template_source=r"\documentclass{article}"), "pdf")

    async def test_runs_pdflatex_when_enabled(self, monkeypatch: pytest.MonkeyPatch) -> None:
        from core.config import settings

        monkeypatch.setattr(settings, "pdf_latex_enabled", True)
        monkeypatch.setattr(settings, "pdf_latex_command", "pdflatex")
        monkeypatch.setattr(
            LatexRenderer, "availability", staticmethod(lambda: (True, None))
        )

        def fake_run(args: list[str], **kwargs: object) -> subprocess.CompletedProcess:
            for arg in args:
                if arg.startswith("-output-directory="):
                    from pathlib import Path

                    out = Path(arg.split("=", 1)[1])
                    (out / "resume.pdf").write_bytes(b"%PDF-1.5 fake-latex")
            return subprocess.CompletedProcess(args, 0, stdout=b"", stderr=b"")

        monkeypatch.setattr(subprocess, "run", fake_run)
        data = LatexRenderer().render(_ctx(template_source=r"\documentclass{article}"), "pdf")
        assert data.startswith(b"%PDF-1.5 fake-latex")


class TestDocxRenderer:
    async def test_builds_docx_from_content(self) -> None:
        data = DocxRenderer().render(_ctx(template_source=""), "docx")
        assert data[:2] == b"PK"

        import io

        from docx import Document

        document = Document(io.BytesIO(data))
        text = "\n".join(p.text for p in document.paragraphs)
        assert "Jane Doe" in text
        assert "Senior Software Engineer" in text
        assert "Led #platform team" in text  # literal text, no templating applied

    async def test_pdf_output_requires_soffice(self, monkeypatch: pytest.MonkeyPatch) -> None:
        from core.config import settings

        monkeypatch.setattr(settings, "pdf_libreoffice_command", "soffice-not-found")
        with pytest.raises(RendererUnavailableError, match="LibreOffice"):
            DocxRenderer().render(_ctx(template_source=""), "pdf")

    async def test_soffice_conversion(self, monkeypatch: pytest.MonkeyPatch) -> None:
        from core.config import settings

        monkeypatch.setattr(settings, "pdf_libreoffice_command", "soffice")
        monkeypatch.setattr(
            "shutil.which", lambda name: "C:/fake/soffice.exe" if name == "soffice" else None
        )

        def fake_run(args: list[str], **kwargs: object) -> subprocess.CompletedProcess:
            from pathlib import Path

            outdir = Path(args[args.index("--outdir") + 1])
            (outdir / "resume.pdf").write_bytes(b"%PDF-1.5 fake-soffice")
            return subprocess.CompletedProcess(args, 0, stdout=b"", stderr=b"")

        monkeypatch.setattr(subprocess, "run", fake_run)
        data = DocxRenderer().render(_ctx(template_source=""), "pdf")
        assert data.startswith(b"%PDF-1.5 fake-soffice")

    async def test_fill_uploaded_template(self) -> None:
        import io

        from docx import Document

        buffer = io.BytesIO()
        source = Document()
        source.add_paragraph("{{ profile.full_name }} — {{ profile.title }}")
        source.add_paragraph("Target: {{ job.title }} at {{ job.company }}")
        source.add_paragraph("{{ content.summary }}")
        source.save(buffer)

        ctx = _ctx(template_source="", template_asset=buffer.getvalue())
        data = DocxRenderer().render(ctx, "docx")

        filled = Document(io.BytesIO(data))
        paragraphs = [p.text for p in filled.paragraphs]
        assert "Jane Doe — Senior Software Engineer" in paragraphs
        assert "Target: Staff Engineer at TechCorp" in paragraphs
        assert "Backend engineer" in "\n".join(paragraphs)


class TestRegistry:
    async def test_capabilities_matrix(self) -> None:
        caps = capabilities()
        by_pair = {(c["format"], c["output"]) for c in caps}
        assert ("html", "pdf") in by_pair
        assert ("typst", "pdf") in by_pair
        assert ("latex", "pdf") in by_pair
        assert ("docx", "docx") in by_pair
        assert ("docx", "pdf") in by_pair
        latex = next(c for c in caps if c["format"] == "latex")
        assert latex["available"] is False
        assert latex["reason"]

    async def test_resolve_engine(self) -> None:
        assert resolve_engine("html", "pdf") == "weasyprint"
        assert resolve_engine("docx", "pdf") == "docx+libreoffice"
        with pytest.raises(ValueError, match="cannot produce"):
            resolve_engine("typst", "docx")

    async def test_get_renderer_unknown(self) -> None:
        with pytest.raises(ValueError, match="Unknown template format"):
            get_renderer("doc")


class TestBuiltinTemplateFiles:
    async def test_reads_builtin_directory(self) -> None:
        from services.template_service import _read_builtin_template_files

        entries = _read_builtin_template_files()
        assert entries is not None
        keys = {e["key"] for e in entries}
        assert {"ats-plain", "classic-serif", "modern-typst", "docx-basic"} <= keys
        for entry in entries:
            if entry["format"] == "docx":
                assert entry["source_text"] is None
            else:
                assert entry["source_text"], f"missing source for {entry['key']}"
            assert entry["name"]

    async def test_seeding_is_insert_only(self) -> None:
        """Existing built-in rows must be left untouched (user edits survive restarts)."""
        from models.rendering import ResumeTemplate
        from services.template_service import TemplateService

        class FakeRepo:
            def __init__(self) -> None:
                self.existing_keys = {"ats-plain", "modern-typst"}

            async def get_by_builtin_key(self, key: str) -> ResumeTemplate | None:
                if key in self.existing_keys:
                    row = ResumeTemplate(
                        builtin_key=key, name=key, format="html", source_text="USER EDIT"
                    )
                    row.id = uuid.uuid5(uuid.NAMESPACE_DNS, key)
                    return row
                return None

        class FakeSession:
            def __init__(self) -> None:
                self.added: list[ResumeTemplate] = []
                self.committed = False

            def add(self, obj: ResumeTemplate) -> None:
                self.added.append(obj)

            async def commit(self) -> None:
                self.committed = True

        service = TemplateService.__new__(TemplateService)
        service._session = FakeSession()
        service._repo = FakeRepo()

        seeded = await service.seed_builtins()

        assert seeded == 2
        added_keys = {t.builtin_key for t in service._session.added}
        assert added_keys == {"classic-serif", "docx-basic"}
        assert service._session.committed is True

    async def test_builtin_typst_template_compiles(self) -> None:
        from services.template_service import SAMPLE_PROFILE, _read_builtin_template_files

        entries = _read_builtin_template_files()
        assert entries is not None
        typst_entry = next(e for e in entries if e["format"] == "typst")
        ctx = RenderContext(
            content=CONTENT,
            profile=dict(SAMPLE_PROFILE),
            template_source=typst_entry["source_text"] or "",
        )
        data = TypstRenderer().render(ctx, "pdf")
        assert data.startswith(b"%PDF")

    async def test_builtin_html_template_render(self, weasyprint_stub: list[str]) -> None:
        from services.template_service import SAMPLE_PROFILE, _read_builtin_template_files

        entries = _read_builtin_template_files()
        assert entries is not None
        for entry in (e for e in entries if e["format"] == "html"):
            ctx = RenderContext(
                content=CONTENT,
                profile=dict(SAMPLE_PROFILE),
                template_source=entry["source_text"] or "",
                styles_text=entry.get("styles_text"),
            )
            data = HtmlRenderer().render(ctx, "pdf")
            assert data.startswith(b"%PDF")


class TestPdfImport:
    async def test_import_sections_and_bullets(self) -> None:
        data = _compile_fixture()
        result = import_pdf(data)

        assert result.name == "Jane Doe"
        assert result.title == "Senior Software Engineer"
        assert result.contact is not None and "jane.doe@example.com" in result.contact
        assert result.content.summary is not None
        assert "Backend engineer" in result.content.summary

        section_names = [s.name for s in result.content.sections]
        assert "Experience" in section_names
        assert "Skills" in section_names

        experience = next(s for s in result.content.sections if s.name == "Experience")
        assert any("platform team" in item.lower() for item in experience.items)

    async def test_import_rejects_empty(self) -> None:
        with pytest.raises(ValueError, match="PDF"):
            import_pdf(b"not a pdf")

    async def test_import_real_typst_output(self) -> None:
        data = TypstRenderer().render(_ctx(), "pdf")
        result = import_pdf(data)
        names = [s.name for s in result.content.sections]
        assert "Experience" in names


def _compile_fixture() -> bytes:
    import tempfile
    from pathlib import Path

    import typst

    with tempfile.TemporaryDirectory() as tmpdir:
        path = Path(tmpdir) / "fixture.typ"
        path.write_text(FIXTURE_TYP, encoding="utf-8")
        return typst.compile(str(path))


class TestEscaping:
    async def test_latex_escape(self) -> None:
        from renderers.base import latex_escape

        assert latex_escape("100% & $5 #1 _x {y}") == (
            r"100\% \& \$5 \#1 \_x \{y\}"
        )

    async def test_typst_escape(self) -> None:
        from renderers.base import typst_escape

        assert typst_escape("#tag $x$ [b]") == r"\#tag \$x\$ \[b\]"
        assert typst_escape("jane@example.com") == r"jane\@example.com"

    async def test_sandbox_blocks_unsafe(self) -> None:
        from renderers.base import build_jinja_env

        env = build_jinja_env()
        rendered = env.from_string(
            "[{{ profile.full_name.__class__ }}][{{ profile.__init__ }}]"
        ).render(profile=PROFILE)
        # ChainableUndefined silences the error, but no unsafe value may leak.
        assert rendered == "[][]"
        assert env.from_string("{{ profile.full_name }}").render(profile=PROFILE) == "Jane Doe"


class TestSandboxedRender:
    async def test_missing_vars_do_not_crash(self, weasyprint_stub: list[str]) -> None:
        ctx = _ctx(template_source="<p>{{ profile.missing.deep }}</p>")
        HtmlRenderer().render(ctx, "pdf")
        assert "<p></p>" in weasyprint_stub[0]
