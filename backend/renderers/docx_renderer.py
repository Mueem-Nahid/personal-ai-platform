from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from docx import Document as DocumentFactory
from docx.document import Document
from docx.shared import Pt

from core.config import settings
from renderers.base import (
    BaseRenderer,
    RenderContext,
    RenderError,
    RendererUnavailableError,
    build_jinja_env,
)


class DocxRenderer(BaseRenderer):
    """Resume content → DOCX via python-docx; DOCX → PDF via LibreOffice headless.

    Two modes:
    - template_asset provided: fill scalar Jinja placeholders ({{ profile.full_name }},
      {{ content.summary }}, {{ job.title }}, ...) inside an uploaded .docx.
    - no asset: build a clean resume document programmatically from the content.
    """

    name = "docx"

    def supported_outputs(self) -> list[str]:
        return ["docx", "pdf"]

    @staticmethod
    def soffice_available() -> bool:
        return shutil.which(settings.pdf_libreoffice_command) is not None

    def render(self, ctx: RenderContext, output_format: str) -> bytes:
        if output_format not in ("docx", "pdf"):
            raise RenderError(f"DOCX renderer cannot produce '{output_format}' output")

        if ctx.template_asset:
            document = _fill_template(ctx.template_asset, ctx)
        else:
            document = _build_document(ctx)

        docx_bytes = _save_document(document)

        if output_format == "docx":
            return docx_bytes
        return _soffice_to_pdf(docx_bytes)


def _placeholder_env() -> Any:
    return build_jinja_env(autoescape=False)


def _iter_paragraphs(document: Document) -> Any:
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph
    for section in document.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                yield paragraph


def _render_paragraph_text(paragraph: Any, context: dict[str, Any]) -> None:
    """Render {{ ... }} placeholders spanning multiple runs in a paragraph.

    Formatting of the first run is preserved for the whole paragraph; block
    constructs ({% for %}) are not supported inside .docx templates.
    """
    original = "".join(run.text for run in paragraph.runs)
    if not original or "{{" not in original:
        return
    rendered = _placeholder_env().from_string(original).render(**context)
    if rendered == original:
        return
    if paragraph.runs:
        paragraph.runs[0].text = rendered
        for run in paragraph.runs[1:]:
            run.text = ""


def _fill_template(asset: bytes, ctx: RenderContext) -> Document:
    import io

    try:
        document = DocumentFactory(io.BytesIO(asset))
    except Exception as e:
        raise RenderError(f"Invalid .docx template: {e}") from e

    context = {
        "profile": ctx.profile or {},
        "content": ctx.content,
        "job": ctx.job or {},
    }
    for paragraph in _iter_paragraphs(document):
        _render_paragraph_text(paragraph, context)
    return document


def _build_document(ctx: RenderContext) -> Document:
    document = DocumentFactory()
    profile = ctx.profile or {}

    name = profile.get("full_name") or "Resume"
    name_heading = document.add_heading(name, level=0)
    for run in name_heading.runs:
        run.font.size = Pt(22)

    contact_parts = [
        str(profile[k]) for k in ("title", "email", "phone", "location") if profile.get(k)
    ]
    if contact_parts:
        document.add_paragraph(" | ".join(contact_parts))

    link_keys = ("github_url", "linkedin_url", "website")
    link_parts = [str(profile[k]) for k in link_keys if profile.get(k)]
    if link_parts:
        document.add_paragraph(" | ".join(link_parts))

    if ctx.content.summary:
        document.add_heading("Summary", level=1)
        document.add_paragraph(ctx.content.summary)

    for section in ctx.content.sections:
        document.add_heading(section.name, level=1)
        for item in section.items:
            document.add_paragraph(item, style="List Bullet")

    return document


def _save_document(document: Document) -> bytes:
    import io

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _soffice_to_pdf(docx_bytes: bytes) -> bytes:
    command = shutil.which(settings.pdf_libreoffice_command)
    if command is None:
        raise RendererUnavailableError(
            f"'{settings.pdf_libreoffice_command}' not found on PATH; "
            "install LibreOffice in the backend image to render DOCX → PDF."
        )

    with tempfile.TemporaryDirectory(prefix="soffice-render-") as tmpdir:
        tmp = Path(tmpdir)
        (tmp / "resume.docx").write_bytes(docx_bytes)
        profile_dir = tmp / "lo-profile"
        profile_dir.mkdir()
        try:
            result = subprocess.run(
                [
                    command,
                    "--headless",
                    "--norestore",
                    f"-env:UserInstallation={profile_dir.resolve().as_uri()}",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(tmp),
                    str(tmp / "resume.docx"),
                ],
                capture_output=True,
                text=True,
                timeout=settings.pdf_render_timeout_seconds,
                check=False,
            )
        except subprocess.TimeoutExpired as e:
            raise RenderError(
                f"LibreOffice timed out after {settings.pdf_render_timeout_seconds}s"
            ) from e

        pdf_path = tmp / "resume.pdf"
        if result.returncode != 0 or not pdf_path.exists():
            output_tail = (result.stdout or result.stderr or "")[-1500:]
            raise RenderError(f"LibreOffice conversion failed:\n{output_tail}")
        return pdf_path.read_bytes()
